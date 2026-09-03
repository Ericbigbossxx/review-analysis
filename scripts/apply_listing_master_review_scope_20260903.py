#!/usr/bin/env python3
"""Apply the 2026-09-03 Review-scope update to config/listing_master.xlsx.

Source of truth for the change: 'THD-Lowe's-Walmart-产品清单 9.3.xlsx' (30 rows,
user-approved on 2026-09-03).  Rules confirmed by the user:
  * New file name/URL wins for every listing present in it.
  * Rows absent from the new file (23 current Review rows) are KEPT in the
    workbook but get monitor_review=FALSE (no longer monitored).  Nothing is
    deleted from the workbook.
  * Every new-file row is an independent listing (no merging by name), with
    monitor_review=TRUE, active=TRUE.
  * WALMART_ORIONX7 is removed from Review monitoring (absent from new file).

Run modes:
  python3 apply_listing_master_review_scope_20260903.py            # apply (writes xlsx)
  python3 apply_listing_master_review_scope_20260903.py --dry-run  # print plan only
  python3 apply_listing_master_review_scope_20260903.py --verify   # validate result
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
MASTER = ROOT / "config" / "listing_master.xlsx"
SOURCE = ROOT / "THD-Lowe's-Walmart-产品清单 9.3.xlsx"

PLATFORM_MAP = {"THD": "THD", "Lowe's": "LOWES", "Walmart": "WALMART"}
PLATFORM_LABEL = {"THD": "Homedepot", "LOWES": "Lowes", "WALMART": "Walmart"}


def s(value) -> str:
    return str(value or "").strip()


def item_id_from_url(url: str) -> str:
    m = re.findall(r"(\d{5,})", url.split("?")[0])
    return m[-1] if m else ""


def product_line(name: str) -> str:
    n = name.lower()
    if "robot" in n:
        return "Robotic Lawn Mower"
    if "snow" in n:
        return "Snow Blower"
    if "mower" in n:
        return "Lawn Mower"
    if "hedge" in n:
        return "Hedge Trimmer"
    if "auger" in n or "hole digger" in n:
        return "Auger"
    if "pole saw" in n:
        return "Pole Saw"
    if "edger" in n:
        return "Edger"
    if "leaf blower" in n or ("backpack" in n and "blower" in n):
        return "Leaf Blower"
    if "trimmer" in n or "brush cutter" in n or "weed" in n:
        return "String Trimmer"
    if "blower" in n:
        return "Leaf Blower"
    return "Yard Care"


def read_new_rows():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb["Sheet1"]
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=4, values_only=True):
        plat, sku, name, url = [s(c) for c in r]
        if not sku:
            continue
        code = PLATFORM_MAP[plat]
        rows.append({
            "platform": code,
            "platform_label": PLATFORM_LABEL[code],
            "record_id": f"{code}_{sku}",
            "internal_sku": sku,
            "product_name": name,
            "listing_url": url,
            "platform_item_id": item_id_from_url(url),
            "product_line": product_line(name),
            "brand": "WILD BADGER POWER" if not name.lower().startswith("sunseeker") and "linkon" not in name.lower() else ("Sunseeker" if name.lower().startswith("sunseeker") else "LinkOn"),
        })
    return rows


def read_master():
    wb = openpyxl.load_workbook(MASTER, data_only=False)
    ws = wb["Listing Master"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [s(h) for h in rows[0]]
    idx = {h: i for i, h in enumerate(hdr)}
    out = []
    for n, vals in enumerate(rows[1:], start=2):
        if not any(s(v) for v in vals):
            continue
        out.append({"row": n, "d": {h: vals[idx[h]] for h in hdr}})
    return wb, ws, hdr, out


def parse_bool(v) -> bool:
    return s(v).upper() in {"TRUE", "1", "YES", "Y"}


def plan():
    new_rows = read_new_rows()
    wb, ws, hdr, master = read_master()

    # current Review-scope keys
    cur_review = {}
    for it in master:
        d = it["d"]
        if parse_bool(d["active"]) and parse_bool(d["monitor_review"]):
            cur_review[(s(d["platform"]).upper(), s(d["internal_sku"]))] = it

    new_by_key = {(r["platform"], r["internal_sku"]): r for r in new_rows}

    actions = []  # dicts describing edits
    # 1) rows in new file that exist as review rows -> update url/name/line/item_id (keep flags)
    # 2) rows in new file not in master at all -> append
    # 3) review rows absent from new file -> monitor_review FALSE
    master_by_key = {}
    for it in master:
        master_by_key.setdefault((s(it["d"]["platform"]).upper(), s(it["d"]["internal_sku"])), it)

    updates = []
    additions = []
    for key, nr in sorted(new_by_key.items()):
        if key in cur_review:
            updates.append((cur_review[key], nr))
        elif key in master_by_key:
            # exists in master but not currently review-monitored -> enable review
            additions.append((master_by_key[key], nr, "enable"))
        else:
            additions.append((None, nr, "append"))

    removals = []
    for key, it in sorted(cur_review.items()):
        if key not in new_by_key:
            removals.append(it)

    return {
        "hdr": hdr, "new_rows": new_rows, "cur_review": cur_review,
        "updates": updates, "additions": additions, "removals": removals,
        "master": master,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--verify", action="store_true")
    args = ap.parse_args()

    p = plan()
    if args.verify:
        verify(p)
        return

    print(f"New-file rows: {len(p['new_rows'])}")
    print(f"Current Review rows in master: {len(p['cur_review'])}")
    print(f"  updates (same SKU, refresh url/name): {len(p['updates'])}")
    print(f"  additions (enable/append): {len(p['additions'])}")
    print(f"  removals (monitor_review -> FALSE): {len(p['removals'])}")
    print()
    print("== UPDATES (keep record_id + monitor_review=TRUE, refresh fields) ==")
    for it, nr in p["updates"]:
        d = it["d"]
        print(f"  row {it['row']:>3} {d['record_id']:<28} url: {s(d['listing_url'])[:60]} -> {nr['listing_url'][:60]}")
    print()
    print("== ADDITIONS ==")
    for it, nr, kind in p["additions"]:
        if it is None:
            print(f"  APPEND  {nr['record_id']:<28} {nr['listing_url'][:70]}")
        else:
            print(f"  ENABLE  row {it['row']:>3} {nr['record_id']:<28} (was monitor_review={s(it['d']['monitor_review'])})")
    print()
    print("== REMOVALS (monitor_review -> FALSE) ==")
    for it in p["removals"]:
        d = it["d"]
        print(f"  row {it['row']:>3} {d['record_id']:<28} {s(d['internal_sku']):<14} {s(d['product_name'])[:45]}")

    if args.dry_run:
        print("\n[dry-run] no changes written.")
        return

    apply(p)


def apply(p):
    wb, ws, hdr, master = p["master"] if False else (None, None, None, None)
    # re-open for write
    wb = openpyxl.load_workbook(MASTER, data_only=False)
    ws = wb["Listing Master"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [s(h) for h in rows[0]]
    idx = {h: i for i, h in enumerate(hdr)}

    def set_cell(row, header, value):
        ws.cell(row=row, column=idx[header] + 1, value=value)

    # updates
    for it, nr in p["updates"]:
        row = it["row"]
        set_cell(row, "platform", nr["platform_label"])
        set_cell(row, "product_line", nr["product_line"])
        set_cell(row, "internal_sku", nr["internal_sku"])
        set_cell(row, "platform_item_id", nr["platform_item_id"])
        set_cell(row, "product_name", nr["product_name"])
        set_cell(row, "listing_url", nr["listing_url"])
        set_cell(row, "brand", nr["brand"])
        set_cell(row, "monitor_review", "TRUE")
        set_cell(row, "active", "TRUE")
        set_cell(row, "notes", s(it["d"].get("notes")) + " | 2026-09-03 scope refresh: URL/product per 9.3 listing file" if s(it["d"].get("notes")) else "2026-09-03 scope refresh: URL/product per 9.3 listing file")
    # additions: enable existing non-review row / append
    max_row = ws.max_row
    for it, nr, kind in p["additions"]:
        if it is None:  # append
            max_row += 1
            row = max_row
            set_cell(row, "record_id", nr["record_id"])
            set_cell(row, "active", "TRUE")
            set_cell(row, "platform", nr["platform_label"])
            set_cell(row, "brand", nr["brand"])
            set_cell(row, "product_line", nr["product_line"])
            set_cell(row, "internal_sku", nr["internal_sku"])
            set_cell(row, "model", None)
            set_cell(row, "platform_item_id", nr["platform_item_id"])
            set_cell(row, "product_name", nr["product_name"])
            set_cell(row, "listing_url", nr["listing_url"])
            set_cell(row, "monitor_listing", "FALSE")
            set_cell(row, "monitor_rank", "FALSE")
            set_cell(row, "monitor_review", "TRUE")
            set_cell(row, "notes", "2026-09-03 added to Review scope from 9.3 listing file under user authorization")
        else:  # enable existing
            row = it["row"]
            set_cell(row, "monitor_review", "TRUE")
            set_cell(row, "active", "TRUE")
            set_cell(row, "platform", nr["platform_label"])
            set_cell(row, "product_line", nr["product_line"])
            set_cell(row, "internal_sku", nr["internal_sku"])
            set_cell(row, "platform_item_id", nr["platform_item_id"])
            set_cell(row, "product_name", nr["product_name"])
            set_cell(row, "listing_url", nr["listing_url"])
            set_cell(row, "brand", nr["brand"])
            set_cell(row, "notes", s(it["d"].get("notes")) + " | 2026-09-03 re-enabled into Review scope from 9.3 file" if s(it["d"].get("notes")) else "2026-09-03 re-enabled into Review scope from 9.3 file")
    # removals
    for it in p["removals"]:
        row = it["row"]
        set_cell(row, "monitor_review", "FALSE")
        notes = s(it["d"].get("notes"))
        if "2026-09-03 review-monitoring ended" not in notes:
            set_cell(row, "notes", (notes + " | " if notes else "") + "2026-09-03 review-monitoring ended (absent from 9.3 listing file)")

    wb.save(MASTER)
    print(f"\nApplied: {len(p['updates'])} updates, {len(p['additions'])} additions, {len(p['removals'])} removals -> {MASTER}")


def verify(p=None):
    if p is None:
        p = plan()
    from collections import Counter
    sys.path.insert(0, str(ROOT))
    from modules.review_tracker.scope import load_review_scope
    scope = load_review_scope(MASTER)
    c = Counter(r["platform_code"] for r in scope)
    print("VERIFY post-change Review scope:", dict(sorted(c.items())), "total", len(scope))
    print("Expected: THD 10 / LOWES 10 / WALMART 10 -> 30")
    ok = len(scope) == 30 and c["THD"] == 10 and c["LOWES"] == 10 and c["WALMART"] == 10
    # every new-file key must be present
    new_rows = read_new_rows()
    keys = {(r["platform"], r["internal_sku"]) for r in new_rows}
    scope_keys = {(r["platform_code"], r["internal_sku"]) for r in scope}
    missing = sorted(keys - scope_keys)
    extra = sorted(scope_keys - keys)
    print("Missing from new-file scope:", missing)
    print("Scope keys not in new file (should be empty, since removals keep row but drop review):", extra)
    if missing or extra:
        ok = False
    print("VERIFY:", "PASS" if ok else "FAIL")


if __name__ == "__main__":
    main()
